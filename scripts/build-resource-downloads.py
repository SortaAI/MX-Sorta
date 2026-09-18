"""Generate original Sorta resources. Dependencies: python-docx, openpyxl, reportlab.
Run from the repo root: python3 scripts/build-resource-downloads.py
"""
from pathlib import Path
from datetime import datetime, time
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.worksheet.datavalidation import DataValidation
from reportlab.pdfgen import canvas
from reportlab.lib.pagesizes import letter
from reportlab.lib.colors import HexColor
OUT=Path(__file__).resolve().parents[1]/'assets/downloads'
OUT.mkdir(exist_ok=True)
fields=['Consultorio / establecimiento','Domicilio del establecimiento','Fecha de registro (dd/mm/aaaa)','Número de expediente / referencia','Nombre completo del paciente','Fecha de nacimiento (dd/mm/aaaa)','Edad a la fecha de registro','Sexo','Domicilio del paciente','Teléfono de contacto','Correo electrónico (si corresponde)','Datos verificados por','Fecha de actualización (dd/mm/aaaa)']
doc=Document();sec=doc.sections[0];sec.top_margin=sec.bottom_margin=Inches(.65)
style=doc.styles['Normal'];style.font.name='Calibri';style.font.size=Pt(10)
doc.add_heading('Ficha de identificación del paciente',0)
doc.add_paragraph('Sorta · Formato administrativo editable · Versión 17/09/2026')
doc.add_paragraph('Adapta los campos al proceso de tu consultorio. Solicita únicamente la información necesaria. Esta ficha no sustituye la historia clínica ni un consentimiento informado.')
table=doc.add_table(rows=0,cols=2);table.style='Table Grid'
for field in fields:
 cells=table.add_row().cells;cells[0].text=field;cells[1].text=' '
doc.add_paragraph('Revisa los datos con el paciente antes de reutilizarlos en otros documentos. Resguarda las fichas según los permisos y procedimientos de tu clínica.')
doc.add_paragraph('mx.getsorta.io/recursos/ficha-identificacion-paciente')
doc.save(OUT/'ficha-identificacion-paciente.docx')

def pdf_start(filename,title,subtitle):
 c=canvas.Canvas(str(OUT/filename),pagesize=letter);c.setTitle(title);c.setAuthor('Equipo de Sorta')
 c.setFillColor(HexColor('#000054'));c.setFont('Helvetica-Bold',20);c.drawString(42,744,title)
 c.setFont('Helvetica',9);c.drawString(42,724,subtitle);return c
c=pdf_start('ficha-identificacion-paciente.pdf','Ficha de identificación del paciente','Sorta · Formato administrativo para imprimir · 17/09/2026')
c.setFont('Helvetica',9)
for i,line in enumerate(['Adapta los campos a tu consultorio y solicita solo la información necesaria.','Esta ficha no sustituye la historia clínica ni el consentimiento informado.']):c.drawString(42,701-i*14,line)
y=652
for field in fields:
 c.setFillColor(HexColor('#66728f'));c.setFont('Helvetica',9);c.drawString(42,y,field)
 c.setStrokeColor(HexColor('#cbd5e7'));c.line(42,y-21,570,y-21);y-=42
c.setFont('Helvetica',8);c.drawString(42,65,'Verifica los datos antes de reutilizarlos. Aplica los procedimientos de resguardo de tu clínica.')
c.drawString(42,48,'mx.getsorta.io/recursos/ficha-identificacion-paciente');c.save()

c=pdf_start('checklist-revision-nom-004.pdf','Revisión documental · NOM-004','Selección orientativa de puntos · No es una auditoría ni certificación · 17/09/2026')
c.setFont('Helvetica',10);c.drawString(42,695,'Consultorio: ________________________   Fecha: ______________')
items=[('Identificación del establecimiento y paciente','5.2 y 5.9'),('Fecha, hora, autor y firma de las notas','5.10'),('Legibilidad y conservación de registros','5.11'),('Contenido correspondiente a la atención prestada','6 a 10'),('Consentimientos: contenido y firmas aplicables','10.1'),('Resguardo, conservación y confidencialidad','5.4 a 5.7')]
y=653
for label,ref in items:
 c.setFont('Helvetica-Bold',11);c.drawString(42,y,label);c.setFont('Helvetica',9);c.drawString(42,y-16,'Referencia NOM-004: '+ref)
 c.drawString(42,y-33,'Estado: Revisado / Pendiente / No aplica     Responsable: __________________')
 c.drawString(42,y-50,'Observaciones: _________________________________________________________');y-=86
c.setFont('Helvetica',8);c.drawString(42,106,'Consulta la norma completa y las disposiciones aplicables a tu establecimiento.')
c.drawString(42,92,'La revisión clínica corresponde al profesional responsable. Autofill no garantiza cumplimiento.')
c.drawString(42,67,'Fuente: https://sidof.segob.gob.mx/notas/docFuente/5272787')
c.linkURL('https://sidof.segob.gob.mx/notas/docFuente/5272787',(42,64,540,78),relative=0)
c.drawString(42,49,'mx.getsorta.io/recursos/llenado-formatos-nom-004');c.save()

wb=Workbook();intro=wb.active;intro.title='Instrucciones';intro.column_dimensions['A'].width=110
instructions=['AGENDA DE CITAS MÉDICAS · SORTA','Plantilla administrativa sin macros · 17/09/2026','1. La hoja Agenda tiene 200 filas en blanco. La hoja Ejemplo usa datos ficticios.','2. Escribe fechas como fechas de Excel y horas de 24 horas: 09:00, 09:30.','3. Duración calcula minutos si la hora final es posterior a la inicial, dentro del mismo día.','4. Elige el estado de la lista y filtra por fecha, profesional o estado.','5. Comprueba manualmente los empalmes. No se detectan citas simultáneas.','6. No hay reservas en línea, mensajes automáticos, contraseña ni control de acceso.','7. Usa Observaciones solo para información administrativa, no diagnósticos.','8. Define quién actualiza el archivo y dónde se conserva la copia vigente.','9. Para superar 200 filas, amplía la tabla y copia fórmulas y validaciones.','10. Si importas en otro editor, revisa fórmulas, listas y formatos después de importar.','Plantilla gratuita: mx.getsorta.io/recursos/agenda-citas-medicas-excel','Conoce la agenda de Sorta: mx.getsorta.io/producto/agenda']
for row,text in enumerate(instructions,1):intro.cell(row,1,text);intro.cell(row,1).alignment=Alignment(wrap_text=True);intro.row_dimensions[row].height=32
headers=['Fecha','Inicio','Fin','Duración (min)','Paciente / referencia','Teléfono','Profesional','Estado','Observaciones administrativas']
for name,example in [('Agenda',False),('Ejemplo',True)]:
 ws=wb.create_sheet(name);ws.append(headers)
 end=4 if example else 201
 for r in range(2,end+1):
  ws.cell(r,4,f'=IF(AND(ISNUMBER(B{r}),ISNUMBER(C{r}),C{r}>B{r}),(C{r}-B{r})*1440,"")')
  for col,fmt in [(1,'dd/mm/yyyy'),(2,'hh:mm'),(3,'hh:mm'),(4,'0'),(6,'@')]:ws.cell(r,col).number_format=fmt
 if example:
  for r,start,status,patient in [(2,9,'Confirmada','DEMO-001 · María López'),(3,10,'Pendiente','DEMO-002 · Luis Pérez'),(4,11,'Cancelada','DEMO-003 · Ana Ruiz')]:
   for col,value in [(1,datetime(2026,9,17)),(2,time(start)),(3,time(start,30)),(5,patient),(7,'Profesional de ejemplo'),(8,status),(9,'Datos ficticios')]:ws.cell(r,col,value)
 dv=DataValidation(type='list',formula1='"Pendiente,Confirmada,Atendida,Cancelada,No asistió"',allow_blank=True);dv.errorTitle='Elige un estado';dv.error='Selecciona una opción de la lista.';dv.showErrorMessage=True;ws.add_data_validation(dv);dv.add(f'H2:H{end}')
 table=Table(displayName='Citas'+name,ref=f'A1:I{end}');table.tableStyleInfo=TableStyleInfo(name='TableStyleMedium2',showRowStripes=True);ws.add_table(table)
 ws.freeze_panes='A2';ws.sheet_view.zoomScale=90
 for col,width in zip('ABCDEFGHI',[15,12,12,18,30,20,24,18,40]):ws.column_dimensions[col].width=width
 ws.print_title_rows='1:1';ws.sheet_properties.pageSetUpPr.fitToPage=True;ws.page_setup.orientation='landscape';ws.page_setup.paperSize=ws.PAPERSIZE_A4;ws.page_setup.fitToWidth=1;ws.page_setup.fitToHeight=0
 for cell in ws[1]:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='000054');cell.alignment=Alignment(wrap_text=True)
 ws.row_dimensions[1].height=32
intro['A1'].font=Font(size=18,bold=True,color='000054')
wb.save(OUT/'agenda-citas-medicas.xlsx')
print('Generated DOCX, XLSX and two PDFs')
