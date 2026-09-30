"""Build original clinic tools. Dependencies: openpyxl, reportlab. No patient data."""
from pathlib import Path
from datetime import datetime, time
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.worksheet.datavalidation import DataValidation
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table as PDFTable, TableStyle
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.colors import HexColor
from reportlab.lib.pagesizes import letter
OUT=Path(__file__).resolve().parents[1]/'assets/downloads'
wb=Workbook();guide=wb.active;guide.title='Instrucciones';guide.column_dimensions['A'].width=105
instructions=['BITÁCORA DE PENDIENTES DE RECEPCIÓN · SORTA','Versión 29/09/2026 · Herramienta administrativa gratuita',
'1. Pendientes incluye 100 filas en blanco; Ejemplo contiene tres casos ficticios.',
'2. Registra una referencia interna, tarea, responsable y siguiente acción concreta.',
'3. Escribe fecha y hora de revisión. Esta hoja no envía avisos ni detecta vencimientos.',
'4. Selecciona Pendiente, En proceso, En espera, Resuelto o Cancelado. Usa los filtros para entregar el turno.',
'5. Conserva una sola versión vigente y define quién la actualiza. No hay controles de acceso configurados.',
'6. No registres diagnósticos, contraseñas ni enlaces privados. Las referencias internas también requieren protección.',
'7. La hoja no sustituye la agenda ni el expediente. Actualiza la herramienta correspondiente al resolver la tarea.',
'8. Para agregar filas, amplía la tabla y copia validaciones. Verifica el archivo si lo importas en otro editor.',
'Guía: https://mx.getsorta.io/recursos/checklist-recepcion-clinica#bitacora']
for i,t in enumerate(instructions,1):guide.cell(i,1,t).alignment=Alignment(wrap_text=True,vertical='center');guide.row_dimensions[i].height=34
headers=['Referencia interna','Tarea administrativa','Responsable','Siguiente acción','Fecha de revisión','Hora de revisión','Estado','Notas de entrega']
examples=[['DEMO-001','Cambio de horario','Turno de tarde','Revisar disponibilidad',datetime(2026,9,29),time(16),'Pendiente','Esperar confirmación antes de cambiar la cita'],['DEMO-002','Solicitud sin confirmar','Recepción','Contactar por el canal acordado',datetime(2026,9,29),time(16,30),'En espera','No asumir cancelación'],['DEMO-003','Formato administrativo incompleto','Recepción','Verificar campo faltante',datetime(2026,9,29),time(17),'En proceso','Sin contenido clínico en esta hoja']]
for name,example in [('Pendientes',False),('Ejemplo',True)]:
 ws=wb.create_sheet(name);ws.append(headers)
 for row in (examples if example else [[None]*8 for _ in range(100)]):ws.append(row)
 end=4 if example else 101
 validation=DataValidation(type='list',formula1='"Pendiente,En proceso,En espera,Resuelto,Cancelado"',allow_blank=True);validation.errorTitle='Selecciona un estado';validation.error='Usa la lista de estados.';validation.showErrorMessage=True;validation.errorStyle='stop';ws.add_data_validation(validation);validation.add(f'G2:G{end}')
 tab=Table(displayName='Tareas'+name,ref=f'A1:H{end}');tab.tableStyleInfo=TableStyleInfo(name='TableStyleMedium2',showRowStripes=True);ws.add_table(tab);ws.freeze_panes='A2'
 for col,width in zip('ABCDEFGH',[22,32,24,36,21,20,19,48]):ws.column_dimensions[col].width=width
 for cell in ws[1]:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='000054');cell.alignment=Alignment(wrap_text=True)
 ws.row_dimensions[1].height=32
 for row in ws.iter_rows(min_row=2,max_row=end):
  for cell in row:cell.alignment=Alignment(wrap_text=True,vertical='top')
  row[4].number_format='dd/mm/yyyy';row[5].number_format='hh:mm'
 ws.print_title_rows='1:1';ws.page_setup.orientation='landscape';ws.page_setup.paperSize=ws.PAPERSIZE_A4;ws.sheet_properties.pageSetUpPr.fitToPage=True;ws.page_setup.fitToWidth=1;ws.page_setup.fitToHeight=0
wb.save(OUT/'bitacora-pendientes-recepcion.xlsx')
styles=getSampleStyleSheet();styles['Title'].textColor=HexColor('#000054');styles['BodyText'].fontSize=9;styles['BodyText'].leading=13
P=lambda t:Paragraph(t,styles['BodyText'])
items=[Paragraph('Checklist administrativa de teleconsulta',styles['Title']),P('Sorta · 29/09/2026 · Preparación de recepción · Sin datos de pacientes'),Spacer(1,12),P('Clínica: __________________  Fecha: __________  Responsable: __________________'),Spacer(1,12)]
rows=[[P('<b>Tarea</b>'),P('<b>Responsable</b>'),P('<b>Estado</b>')]]
for text in ['ANTES DE LA CITA','Confirmar fecha, hora y zona horaria','Verificar plataforma e instrucciones de acceso','Definir contacto de ayuda y canal autorizado','Revisar formularios administrativos pendientes','AL CONECTARSE','Verificar la cita según el proceso de la clínica','Comprobar acceso, audio y cámara si solicita ayuda','Avisar al profesional si existe una incidencia','SI FALLA LA CONEXIÓN','Aplicar instrucciones de soporte de la clínica','Confirmar con el profesional el siguiente paso','Comunicar la decisión y actualizar la agenda','AL CERRAR','Confirmar el estado administrativo con el equipo','Entregar pendientes con responsable y revisión']:
 rows.append([P('<b>'+text+'</b>' if text.isupper() else text),P(''),P('')])
table=PDFTable(rows,colWidths=[330,115,95]);table.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),HexColor('#e8ecff')),('GRID',(0,0),(-1,-1),.4,HexColor('#dce4f2')),('VALIGN',(0,0),(-1,-1),'TOP'),('TOPPADDING',(0,0),(-1,-1),7),('BOTTOMPADDING',(0,0),(-1,-1),7)]));items += [table,Spacer(1,12),P('Estado sugerido: pendiente / en proceso / hecho. No es un protocolo clínico ni un consentimiento informado. La pertinencia de la teleconsulta y las decisiones clínicas corresponden al profesional. Sigue el procedimiento de la clínica ante necesidades urgentes.'),Spacer(1,8),P('No anotes diagnósticos, contraseñas ni enlaces privados en esta hoja.'),P('mx.getsorta.io/recursos/checklist-teleconsulta')]
SimpleDocTemplate(str(OUT/'checklist-teleconsulta.pdf'),pagesize=letter,rightMargin=36,leftMargin=36,topMargin=28,bottomMargin=28,title='Checklist administrativa de teleconsulta | Sorta',author='Equipo de Sorta').build(items)
check=load_workbook(OUT/'bitacora-pendientes-recepcion.xlsx');assert check.sheetnames==['Instrucciones','Pendientes','Ejemplo'];assert check['Pendientes'].max_row==101;assert check['Ejemplo']['A2'].value=='DEMO-001'
print('Generated and checked reception workbook and teleconsultation PDF')
